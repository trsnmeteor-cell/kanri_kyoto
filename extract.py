import openpyxl, json, datetime
wb = openpyxl.load_workbook('sheet.xlsx', data_only=True)

s = wb['店舗']
store_by_code = {}
for r in range(7, s.max_row+1):
    code = s.cell(row=r, column=6).value
    name = s.cell(row=r, column=7).value
    area = s.cell(row=r, column=5).value
    if code is not None and name:
        try: code = int(code)
        except: pass
        store_by_code[code] = {'name': str(name), 'area': area}

def tds(v): return int(v.total_seconds()) if isinstance(v, datetime.timedelta) else None
def hhmm(sec):
    if sec is None: return None
    return f"{sec//3600}:{(sec%3600)//60:02d}"

ws = wb['超過']
chouka=[]
for r in range(11, ws.max_row+1):
    name = ws.cell(row=r, column=4).value
    if not name: continue
    code = ws.cell(row=r, column=2).value
    try: code=int(code)
    except: pass
    ov = tds(ws.cell(row=r, column=5).value); cnt = ws.cell(row=r, column=6).value; fut = tds(ws.cell(row=r, column=7).value)
    chouka.append({'code':code,'store':store_by_code.get(code,{}).get('name'),'area':store_by_code.get(code,{}).get('area'),
        'emp':str(ws.cell(row=r,column=3).value),'name':str(name),'ovSec':ov,'ov':hhmm(ov),'cnt':cnt,
        'cntNum':int(str(cnt).split('/')[0]) if cnt else 0,'futSec':fut,'futH':round(fut/3600,1) if fut else None})
chouka.sort(key=lambda x:-(x['ovSec'] or 0))

CATS=['その他','飲食','料金','クレンリネス','機械','接客','-']
g = wb['ご意見']
def block(period_row, start, end):
    ps = g.cell(row=period_row, column=34).value
    pe = g.cell(row=period_row, column=36).value
    stores=[]
    for r in range(start, end+1):
        nm = g.cell(row=r, column=33).value
        if not nm: continue
        vals=[g.cell(row=r,column=c).value or 0 for c in range(34,41)]
        vals=[int(v) if isinstance(v,(int,float)) else 0 for v in vals]
        stores.append({'name':str(nm).strip(),'v':vals,'total':sum(vals)})
    return {'key':ps.strftime('%Y-%m'),'label':f"{ps.year}年{ps.month}月",
            'start':ps.strftime('%Y/%m/%d'),'end':pe.strftime('%Y/%m/%d'),
            'stores':stores,'catTotals':[sum(st['v'][i] for st in stores) for i in range(7)],
            'grand':sum(st['total'] for st in stores)}

months=[block(3,6,21), block(24,27,42)]
months.sort(key=lambda m:m['key'], reverse=True)

out={'generated':'2026-10-04','cats':CATS,'chouka':chouka,'months':months}


# --- 店舗チェック（店舗ﾁｪｯｸ / 店舗ﾁｪｯｸd）---
from openpyxl.utils import column_index_from_string
chk = wb['店舗ﾁｪｯｸ']; dchk = wb['店舗ﾁｪｯｸd']
th = {'high': int(chk.cell(row=4, column=9).value or 3), 'mid': int(chk.cell(row=4, column=10).value or 2)}
sc_stores = []
for r in range(5, 31):
    code = chk.cell(row=r, column=3).value; name = chk.cell(row=r, column=4).value
    tcol = chk.cell(row=r, column=5).value
    if code is None or not name or not tcol: continue
    si = column_index_from_string(str(tcol).strip()) - 5
    items = []
    for rr in range(5, dchk.max_row + 1):
        cat = dchk.cell(row=rr, column=si + 2).value; item = dchk.cell(row=rr, column=si + 3).value
        cnt = dchk.cell(row=rr, column=si + 6).value
        if cnt is None or (not cat and not item): continue
        try: cnt = int(round(float(cnt)))
        except: continue
        items.append({'type': str(dchk.cell(row=rr, column=si + 1).value or '').strip(),
                      'cat': str(cat or '').strip(), 'item': str(item or '').strip(), 'cnt': cnt})
    sc_stores.append({'code': int(code), 'name': str(name).strip(), 'items': items,
                      'total': sum(i['cnt'] for i in items), 'max': max([i['cnt'] for i in items], default=0)})
out['storeCheck'] = {'thresholds': th, 'stores': sc_stores}

json.dump(out, open('data.json','w'), ensure_ascii=False, indent=1)
print('months:', [(m['key'],m['label'],m['grand']) for m in months])
print('chouka:', len(chouka), 'cats', CATS)
